# Builds the Excel versions of the free jewellery formats into public/formats/ (one .xlsx per format,
# named after its page). The sums match src/lib/jewelleryFormats.js. Run after changing a format:
#   python scripts/formats-xlsx.py
# Needs openpyxl (pip install openpyxl). The .xlsx files are committed, so the site build does not run this.
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

OUT = Path(__file__).resolve().parent.parent / "public" / "formats"
INPUT = PatternFill("solid", fgColor="FFF8DC")
HEAD = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="555555")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
ROWS = 10


def sheet(title, widths):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    return wb, ws


def last_col(ws):
    return chr(64 + ws.max_column)


def shop_header(ws, cols, tax_label):
    end = chr(64 + cols)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = "YOUR SHOP NAME"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = "Address · Phone"
    ws.merge_cells(f"A3:{end}3")
    ws["A3"] = f"{tax_label}: "
    for c in ("A1", "A2", "A3"):
        ws[c].fill = INPUT
        ws[c].alignment = Alignment(horizontal="center")


def title(ws, row, cols, text):
    end = chr(64 + cols)
    ws.merge_cells(f"A{row}:{end}{row}")
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = Font(bold=True, size=13)
    ws[f"A{row}"].alignment = Alignment(horizontal="center")


def fields(ws, row, pairs):
    """[(label, col_letter_of_label)] laid out as label | input over two columns."""
    for r, items in enumerate(pairs):
        for label, col in items:
            c = ws[f"{col}{row + r}"]
            c.value = label
            c.font = BOLD
            inp = ws.cell(row=row + r, column=c.column + 1)
            inp.fill = INPUT
            inp.border = BOX
    return row + len(pairs)


def table(ws, row, heads):
    for i, h in enumerate(heads):
        c = ws.cell(row=row, column=i + 1, value=h)
        c.font = BOLD
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    return row + 1


def grid(ws, first, last, cols, inputs):
    for r in range(first, last + 1):
        for col in range(1, cols + 1):
            c = ws.cell(row=r, column=col)
            c.border = BOX
            if chr(64 + col) in inputs:
                c.fill = INPUT


def total(ws, row, label_col, label, formula, fmt="#,##0.00", inp=False, bold=False):
    ws[f"{label_col}{row}"] = label
    ws[f"{label_col}{row}"].font = Font(bold=bold)
    vcol = chr(ord(label_col) + 1)
    c = ws[f"{vcol}{row}"]
    c.value = formula
    c.number_format = fmt
    c.border = BOX
    if inp:
        c.fill = INPUT
    if bold:
        c.font = BOLD
    return row + 1


def note(ws, row, cols, text, italic=True):
    end = chr(64 + cols)
    ws.merge_cells(f"A{row}:{end}{row}")
    ws[f"A{row}"] = text
    ws[f"A{row}"].font = Font(italic=italic, size=9)
    ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 26
    return row + 1


def signatures(ws, row, left, right, cols):
    ws[f"A{row + 2}"] = "______________________"
    ws[f"A{row + 3}"] = left
    rc = chr(64 + cols - 2)
    ws[f"{rc}{row + 2}"] = "______________________"
    ws[f"{rc}{row + 3}"] = right
    return row + 4


FOOT = "Free format from www.infovion.in/jewellery-bill-formats. Yellow cells: type here. White cells calculate. Confirm tax rates with your accountant."


def bill(estimate=False):
    cols = 13 if not estimate else 12
    heads = ["Sr.", "Description", "HSN", "HUID", "Purity", "Gross wt (g)", "Net wt (g)", "Rate / g (₹)", "Making / g (₹)", "Stone / other (₹)", "Metal value", "Making", "Amount (₹)"]
    if estimate:
        heads.remove("HSN")
    widths = [5, 22, 8, 9, 8, 10, 10, 11, 11, 11, 12, 11, 13] if not estimate else [5, 24, 9, 8, 10, 10, 11, 11, 11, 12, 11, 13]
    wb, ws = sheet("Estimate" if estimate else "Tax invoice", widths)
    shop_header(ws, cols, "GSTIN")
    title(ws, 5, cols, "ESTIMATE" if estimate else "TAX INVOICE")
    if estimate:
        fields(ws, 7, [[("Customer", "A"), ("Estimate no.", "H")], [("Mobile", "A"), ("Date", "H")], [("Address", "A"), ("Include GST (Y/N)", "H")]])
        ws["I9"] = "Y"
    else:
        fields(ws, 7, [[("Customer", "A"), ("Invoice no.", "H")], [("Address", "A"), ("Date", "H")], [("Mobile", "A"), ("Place of supply", "H")], [("GSTIN", "A"), ("Inter-state (Y/N)", "H")], [("PAN", "A"), ("Payment", "H")]])
        ws["I10"] = "N"
    top = 12 if not estimate else 11
    first = table(ws, top, heads)
    last = first + ROWS - 1
    # column letters shift by one in the estimate (no HSN)
    L = (lambda c: c) if not estimate else (lambda c: chr(ord(c) - 1) if c > "C" else c)
    G, H, I, J, K, Lm, M = L("G"), L("H"), L("I"), L("J"), L("K"), L("L"), L("M")
    for r in range(first, last + 1):
        ws[f"A{r}"] = r - first + 1
        ws[f"{K}{r}"] = f'=IF({G}{r}="","",ROUND({G}{r}*{H}{r},2))'
        ws[f"{Lm}{r}"] = f'=IF({G}{r}="","",ROUND({G}{r}*{I}{r},2))'
        ws[f"{M}{r}"] = f'=IF({G}{r}="","",{K}{r}+{Lm}{r}+N({J}{r}))'
        for c in (L("F"), G):
            ws[f"{c}{r}"].number_format = "0.000"
        for c in (H, I, J, K, Lm, M):
            ws[f"{c}{r}"].number_format = "#,##0.00"
    inputs = {"B", "C", "D", "E", "F", "G", "H", "I", "J"} if not estimate else {"B", "C", "D", "E", "F", "G", "H", "I"}
    grid(ws, first, last, cols, inputs)
    if not estimate:
        ws[f"C{first}"] = "7113"
    lc = chr(ord(M) - 1)
    r = last + 2
    r = total(ws, r, lc, "Total", f"=SUM({M}{first}:{M}{last})")
    r = total(ws, r, lc, "Less: discount", 0, inp=True)
    taxable = r
    r = total(ws, r, lc, "Taxable value", f"={M}{r - 2}-{M}{r - 1}")
    if estimate:
        r = total(ws, r, lc, "Estimated GST 3%", f'=IF(UPPER(I9)="Y",ROUND({M}{taxable}*3%,2),0)')
        r = total(ws, r, lc, "Estimated total", f"={M}{taxable}+{M}{r - 1}", bold=True)
        r = note(ws, r + 1, cols, "This is an estimate, not a tax invoice. The final bill is made at the rate on the day of purchase.")
    else:
        r = total(ws, r, lc, "CGST 1.5%", f'=IF(UPPER(I10)="Y",0,ROUND({M}{taxable}*1.5%,2))')
        r = total(ws, r, lc, "SGST 1.5%", f'=IF(UPPER(I10)="Y",0,ROUND({M}{taxable}*3%,2)-{M}{r - 1})')
        r = total(ws, r, lc, "IGST 3%", f'=IF(UPPER(I10)="Y",ROUND({M}{taxable}*3%,2),0)')
        inv = r
        r = total(ws, r, lc, "Invoice total", f"={M}{taxable}+{M}{r - 3}+{M}{r - 2}+{M}{r - 1}", bold=True)
        r = total(ws, r, lc, "Less: old gold", 0, inp=True)
        r = total(ws, r, lc, "Less: advance", 0, inp=True)
        r = total(ws, r, lc, "Round off", f"=ROUND({M}{inv}-{M}{r - 2}-{M}{r - 1},0)-({M}{inv}-{M}{r - 2}-{M}{r - 1})")
        r = total(ws, r, lc, "Balance payable (₹)", f"=ROUND({M}{inv}-{M}{r - 3}-{M}{r - 2},0)", "#,##0", bold=True)
        ws[f"A{last + 2}"] = "Old gold taken:"
        ws[f"B{last + 2}"].fill = INPUT
        ws[f"A{last + 4}"] = "Amount in words:"
        ws[f"B{last + 4}"].fill = INPUT
        r = note(ws, r + 1, cols, "Certified that the particulars given above are true and correct. Check the HUID of hallmarked jewellery in the BIS CARE app.")
    r = signatures(ws, r, "Customer's signature", "Authorised signatory", cols)
    note(ws, r + 1, cols, FOOT)
    return wb


def girvi():
    cols = 6
    wb, ws = sheet("Girvi receipt", [6, 30, 8, 13, 13, 12])
    shop_header(ws, cols, "Money-lending licence no.")
    title(ws, 5, cols, "GIRVI (PLEDGE) RECEIPT")
    fields(ws, 7, [[("Pledger", "A"), ("Girvi no.", "D")], [("S/o, D/o or W/o", "A"), ("Date", "D")], [("Address", "A"), ("Due date", "D")], [("Mobile", "A"), ("Loan (₹)", "D")], [("ID proof", "A"), ("Interest % / month", "D")]])
    ws["D12"] = "Interest / month (₹)"
    ws["D12"].font = BOLD
    ws["E12"] = "=ROUND(E10*E11/100,2)"
    ws["E12"].number_format = "#,##0.00"
    first = table(ws, 14, ["Sr.", "Ornament", "Pcs", "Gross wt (g)", "Net wt (g)", "Purity"])
    last = first + 7
    for r in range(first, last + 1):
        ws[f"A{r}"] = r - first + 1
        ws[f"D{r}"].number_format = ws[f"E{r}"].number_format = "0.000"
    grid(ws, first, last, cols, {"B", "C", "D", "E", "F"})
    ws[f"D{last + 1}"] = "Total net wt"
    ws[f"E{last + 1}"] = f"=SUM(E{first}:E{last})"
    ws[f"E{last + 1}"].number_format = "0.000"
    r = last + 3
    ws[f"A{r}"] = "Loan in words:"
    ws[f"B{r}"].fill = INPUT
    r += 2
    ws[f"A{r}"] = "Terms"
    ws[f"A{r}"].font = BOLD
    terms = [
        "1. Interest is simple interest at the rate above on the amount outstanding, counted by days, 30 days being one month.",
        "2. Interest for at least 15 days is charged, even if the ornaments are released earlier.",
        "3. The ornaments are returned only to the pledger, or a person they authorise in writing, against this receipt, on payment of the loan and interest.",
        "4. Weights and purity were checked at the counter in the pledger's presence.",
        "5. If the loan is not repaid by the due date, the shop will give notice and act only as the state money-lending law allows.",
    ]
    for t in terms:
        r = note(ws, r + 1, cols, t, italic=False) - 1
    r = signatures(ws, r + 1, "Pledger's signature", "Authorised signatory", cols)
    r += 1
    ws.merge_cells(f"A{r}:F{r}")
    ws[f"A{r}"] = "- - - - - - - - - - - - - - - -  RELEASE  - - - - - - - - - - - - - - - -"
    ws[f"A{r}"].alignment = Alignment(horizontal="center")
    r = fields(ws, r + 1, [[("Release date", "A"), ("Principal (₹)", "D")], [("Interest (₹)", "A"), ("Total received (₹)", "D")]])
    ws[f"E{r - 1}"] = f"=N(E{r - 2})+N(B{r - 1})"
    r = note(ws, r, cols, "I have received my ornaments back in full.", italic=False)
    r = signatures(ws, r, "Pledger's signature", "Authorised signatory", cols)
    note(ws, r + 1, cols, FOOT)
    return wb


def old_gold():
    cols = 8
    wb, ws = sheet("Old gold voucher", [5, 24, 11, 11, 11, 9, 11, 13])
    shop_header(ws, cols, "GSTIN")
    title(ws, 5, cols, "OLD GOLD PURCHASE VOUCHER")
    fields(ws, 7, [[("Seller", "A"), ("Voucher no.", "F")], [("Address", "A"), ("Date", "F")], [("Mobile", "A"), ("Payment", "F")], [("ID proof", "A"), ("24K rate / g (₹)", "F")], [("PAN", "A"), ("Deduction %", "F")]])
    first = table(ws, 13, ["Sr.", "Description", "Gross wt (g)", "Less wt (g)", "Net wt (g)", "Touch %", "Fine wt (g)", "Value (₹)"])
    last = first + ROWS - 1
    for r in range(first, last + 1):
        ws[f"A{r}"] = r - first + 1
        ws[f"E{r}"] = f'=IF(C{r}="","",ROUND(C{r}-N(D{r}),3))'
        ws[f"G{r}"] = f'=IF(C{r}="","",ROUND(E{r}*F{r}/100,3))'
        ws[f"H{r}"] = f'=IF(C{r}="","",ROUND(G{r}*$G$10,2))'
        for c in "CDEG":
            ws[f"{c}{r}"].number_format = "0.000"
        ws[f"H{r}"].number_format = "#,##0.00"
    grid(ws, first, last, cols, {"B", "C", "D", "F"})
    r = last + 2
    r = total(ws, r, "G", "Fine wt (g)", f"=SUM(G{first}:G{last})", "0.000")
    r = total(ws, r, "G", "Value (₹)", f"=SUM(H{first}:H{last})")
    r = total(ws, r, "G", "Less: deduction", f"=ROUND(H{r - 1}*N(G11)/100,2)")
    r = total(ws, r, "G", "Round off", f"=ROUND(H{r - 2}-H{r - 1},0)-(H{r - 2}-H{r - 1})")
    r = total(ws, r, "G", "Amount paid (₹)", f"=ROUND(H{r - 3}-H{r - 2},0)", "#,##0", bold=True)
    ws[f"A{last + 2}"] = "Amount in words:"
    ws[f"B{last + 2}"].fill = INPUT
    r = note(ws, r + 1, cols, "I declare that I am the lawful owner of the gold sold above, that it is not stolen or pledged anywhere, and that I have received the amount above in full.", italic=False)
    r = signatures(ws, r, "Seller's signature", "Authorised signatory", cols)
    note(ws, r + 1, cols, FOOT)
    return wb


def karigar():
    cols = 7
    wb, ws = sheet("Karigar slip", [5, 26, 8, 12, 9, 10, 13])
    shop_header(ws, cols, "GSTIN")
    title(ws, 5, cols, "KARIGAR ISSUE AND RECEIPT SLIP")
    fields(ws, 7, [[("Karigar", "A"), ("Slip no.", "E")], [("Mobile", "A"), ("Date", "E")], [("Order / design", "A"), ("Due date", "E")], [("Work to be made", "A"), ("Labour / g (₹)", "E")]])
    ws["A12"] = "Metal issued"
    ws["A12"].font = BOLD
    first = table(ws, 13, ["Sr.", "Description", "", "Weight (g)", "Touch %", "", "Fine wt (g)"])
    last = first + 4
    for r in range(first, last + 1):
        ws[f"A{r}"] = r - first + 1
        ws[f"G{r}"] = f'=IF(D{r}="","",ROUND(D{r}*E{r}/100,3))'
        ws[f"D{r}"].number_format = ws[f"G{r}"].number_format = "0.000"
    grid(ws, first, last, cols, {"B", "D", "E"})
    issued = last + 1
    ws[f"F{issued}"] = "Fine issued"
    ws[f"G{issued}"] = f"=SUM(G{first}:G{last})"
    ws[f"G{issued}"].number_format = "0.000"
    ws[f"A{issued + 2}"] = "Work received"
    ws[f"A{issued + 2}"].font = BOLD
    f2 = table(ws, issued + 3, ["Sr.", "Description", "Pcs", "Weight (g)", "Touch %", "Wastage %", "Fine credited (g)"])
    l2 = f2 + 4
    for r in range(f2, l2 + 1):
        ws[f"A{r}"] = r - f2 + 1
        ws[f"G{r}"] = f'=IF(D{r}="","",ROUND(D{r}*E{r}/100*(1+N(F{r})/100),3))'
        ws[f"D{r}"].number_format = ws[f"G{r}"].number_format = "0.000"
    grid(ws, f2, l2, cols, {"B", "C", "D", "E", "F"})
    r = l2 + 2
    r = total(ws, r, "F", "Fine issued (g)", f"=G{issued}", "0.000")
    r = total(ws, r, "F", "Fine credited (g)", f"=SUM(G{f2}:G{l2})", "0.000")
    r = total(ws, r, "F", "Fine balance with karigar (g)", f"=G{r - 2}-G{r - 1}", "0.000", bold=True)
    r = total(ws, r, "F", "Labour payable (₹)", f"=ROUND(SUM(D{f2}:D{l2})*N(F10),2)")
    r = note(ws, r + 1, cols, "Fine credited = weight × touch ÷ 100, plus the agreed wastage as a % of it.")
    r = signatures(ws, r, "Karigar's signature", "Authorised signatory", cols)
    note(ws, r + 1, cols, FOOT)
    return wb


def uae():
    cols = 9
    wb, ws = sheet("Tax invoice (UAE)", [5, 22, 7, 10, 11, 11, 12, 7, 15])
    shop_header(ws, cols, "TRN")
    title(ws, 5, cols, "TAX INVOICE  ·  فاتورة ضريبية")
    fields(ws, 7, [[("Customer", "A"), ("Invoice no.", "G")], [("Address", "A"), ("Invoice date", "G")], [("Mobile", "A"), ("Date of supply", "G")], [("Customer TRN", "A"), ("Reverse charge (Y/N)", "G")]])
    ws["H10"] = "N"
    first = table(ws, 12, ["Sr.", "Description", "Karat", "Net wt (g)", "Rate / g (AED)", "Making / g (AED)", "Stones / other (AED)", "VAT %", "Amount excl. VAT (AED)"])
    last = first + ROWS - 1
    for r in range(first, last + 1):
        ws[f"A{r}"] = r - first + 1
        ws[f"H{r}"] = f'=IF(D{r}="","",IF(UPPER($H$10)="Y","RC","5%"))'
        ws[f"I{r}"] = f'=IF(D{r}="","",ROUND(D{r}*E{r},2)+ROUND(D{r}*N(F{r}),2)+N(G{r}))'
        ws[f"D{r}"].number_format = "0.000"
        for c in "EFGI":
            ws[f"{c}{r}"].number_format = "#,##0.00"
    grid(ws, first, last, cols, {"B", "C", "D", "E", "F", "G"})
    r = last + 2
    r = total(ws, r, "H", "Total excl. VAT", f"=SUM(I{first}:I{last})")
    r = total(ws, r, "H", "Less: discount", 0, inp=True)
    tx = r
    r = total(ws, r, "H", "Taxable amount", f"=I{r - 2}-I{r - 1}")
    r = total(ws, r, "H", "VAT 5%", f'=IF(UPPER(H10)="Y",0,ROUND(I{tx}*5%,2))')
    r = total(ws, r, "H", "Total incl. VAT (AED)", f"=I{tx}+I{r - 1}", bold=True)
    r = total(ws, r, "H", "Less: old gold", 0, inp=True)
    r = total(ws, r, "H", "Balance payable (AED)", f"=I{r - 2}-I{r - 1}", bold=True)
    ws[f"A{last + 2}"] = "Amount in words:"
    ws[f"B{last + 2}"].fill = INPUT
    r = note(ws, r + 1, cols, "Amounts in UAE dirhams. VAT at 5% on the full value of the jewellery: metal, making and stones. If Reverse charge is Y: reverse charge applies, and the recipient is liable to account for the VAT (Cabinet Decision No. 127 of 2024).")
    r = signatures(ws, r, "Customer's signature", "Authorised signatory", cols)
    note(ws, r + 1, cols, FOOT)
    return wb


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    books = {
        "jewellery-bill-format": bill(),
        "jewellery-estimate-format": bill(estimate=True),
        "girvi-receipt-format": girvi(),
        "old-gold-purchase-voucher-format": old_gold(),
        "karigar-issue-slip-format": karigar(),
        "uae-jewellery-tax-invoice-format": uae(),
    }
    for name, wb in books.items():
        wb.properties.creator = "Infovion Technologies"
        wb.properties.title = name.replace("-", " ").capitalize()
        wb.save(OUT / f"{name}.xlsx")
        print("wrote", OUT / f"{name}.xlsx")
